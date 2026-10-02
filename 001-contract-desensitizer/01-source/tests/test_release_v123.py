"""v1.2.3 的识别边界与本机历史记录安全回归。"""
import json
import os
import sys
import tempfile
import threading
import unittest
import urllib.error
import urllib.request
from docx import Document

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import contract_sensitive_detector as D  # noqa: E402
import contract_app_server as S  # noqa: E402


class ReleaseV123Tests(unittest.TestCase):
    def setUp(self):
        self.engine = D.DetectorEngine(D.Settings(min_confidence=0.5, min_severity="low"))

    def hits(self, text, kind):
        doc = D.Document(source="sample.txt", file_type="txt", blocks=[D.Block(text=text)])
        return [hit.value for hit in self.engine.scan_document(doc) if hit.type_id == kind]

    def test_organization_name_does_not_swallow_narrative(self):
        values = self.hits("为进一步规范上海分公司反洗钱工作。", "ORG_NAME")
        self.assertTrue(all("进一步规范" not in value for value in values))
        values = self.hits("上海和远系统集成有限公司北京分公司承担实施。", "ORG_NAME")
        self.assertIn("上海和远系统集成有限公司北京分公司", values)

    def test_quoted_address_keeps_full_value(self):
        values = self.hits("联系地址为“上海市浦东新区示例小区”", "ADDRESS")
        self.assertIn("上海市浦东新区示例小区", values)
        self.assertNotIn("“上海市浦东新区示例小区”", values)

    def test_history_persists_and_delete_checks_request_origin(self):
        with tempfile.TemporaryDirectory() as temp:
            sid = "a1b2c3d4e5f6"
            out = os.path.join(temp, sid, "out")
            os.makedirs(out)
            with open(os.path.join(out, "sample_脱敏.txt"), "w", encoding="utf-8") as file:
                file.write("[自然人姓名#1]")
            server = S.start_app_server("127.0.0.1", 0, temp)
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            base = "http://127.0.0.1:%d" % server.server_port
            try:
                with urllib.request.urlopen(base + "/history", timeout=3) as response:
                    self.assertIn(sid, response.read().decode("utf-8"))
                body = json.dumps({"sid": sid}).encode("utf-8")
                foreign = urllib.request.Request(
                    base + "/api/history/delete", data=body, method="POST",
                    headers={"Content-Type": "application/json", "Origin": "https://other.example"})
                with self.assertRaises(urllib.error.HTTPError) as error:
                    urllib.request.urlopen(foreign, timeout=3)
                self.assertEqual(error.exception.code, 403)
                self.assertTrue(os.path.isdir(os.path.join(temp, sid)))
                local = urllib.request.Request(
                    base + "/api/history/delete", data=body, method="POST",
                    headers={"Content-Type": "application/json", "Origin": base})
                with urllib.request.urlopen(local, timeout=3) as response:
                    self.assertTrue(json.load(response)["ok"])
                self.assertFalse(os.path.exists(os.path.join(temp, sid)))
            finally:
                server.shutdown()
                server.server_close()

    def test_restore_uses_mapping_created_after_manual_review(self):
        with tempfile.TemporaryDirectory() as temp:
            sid = "a1b2c3d4e5f6"
            sess = os.path.join(temp, sid)
            out = os.path.join(sess, "out")
            os.makedirs(out)
            source = os.path.join(sess, "sample.docx")
            redacted = os.path.join(out, "sample_脱敏.docx")
            original = Document()
            original.add_paragraph("项目代号 Project Aurora")
            original.save(source)
            masked = Document()
            masked.add_paragraph("项目代号 [项目代号#1]")
            masked.save(redacted)
            upload_mapping = source + ".mapping.json"
            with open(upload_mapping, "w", encoding="utf-8") as file:
                json.dump({"schema": "contract-mapping/v1", "items": []}, file)
            with open(redacted + ".mapping.json", "w", encoding="utf-8") as file:
                json.dump({"schema": "contract-mapping/v1", "items": [
                    {"placeholder": "[项目代号#1]", "original": "Project Aurora", "start": 5}
                ]}, file)
            server = S.start_app_server("127.0.0.1", 0, temp)
            server.RequestHandlerClass.state.new_session(sid, {
                "input_path": source, "mode": "redact", "sibling_mapping": upload_mapping,
            })
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            try:
                req = urllib.request.Request(
                    "http://127.0.0.1:%d/api/restore/%s" % (server.server_port, sid),
                    data=b"{}", method="POST", headers={"Content-Type": "application/json"})
                with urllib.request.urlopen(req, timeout=3) as response:
                    result = json.load(response)
                self.assertTrue(result["ok"])
                restored = Document(os.path.join(out, "sample_还原.docx"))
                self.assertEqual(restored.paragraphs[0].text, "项目代号 Project Aurora")
            finally:
                server.shutdown()
                server.server_close()


if __name__ == "__main__":
    unittest.main()

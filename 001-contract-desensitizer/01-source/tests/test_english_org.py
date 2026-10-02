#!/usr/bin/env python3
"""英文合同主体及紧邻简称的定向回归。"""

import os
import sys
import tempfile
import unittest
from zipfile import ZipFile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document

import contract_redactor as R
import contract_sensitive_detector as D
from contract_app_server import _run_detector


SAMPLE = (
    '本条款清单由包括【Nova Technologies Limited】（“NOVA”）在内的卖方主体签订。\n'
    'NOVA 应按约定交付。\n'
    'This Term Sheet is between [Nova Technologies Limited] and Zenith Labs Limited.'
)


class EnglishOrgTest(unittest.TestCase):
    def test_full_name_alias_and_restore(self):
        with tempfile.TemporaryDirectory() as tmp:
            src = os.path.join(tmp, "sample.docx")
            doc = Document()
            first, *rest = SAMPLE.splitlines()
            before, after = first.split("Nova Technologies Limited")
            paragraph = doc.add_paragraph()
            paragraph.add_run(before)
            paragraph.add_run("Nova Technologies Limited")  # 灰底字段常是独立 run
            paragraph.add_run(after)
            for line in rest:
                doc.add_paragraph(line)
            doc.save(src)

            report = _run_detector(src)["report"]
            orgs = {item["value"]: item for item in report["items"]
                    if item["type"] == "ORG_NAME"}
            self.assertIn("Nova Technologies Limited", orgs)
            self.assertIn("NOVA", orgs)
            self.assertIn("Zenith Labs Limited", orgs)

            numberer = D.EntityNumberer(report["items"])
            self.assertEqual(numberer.number("ORG_NAME", "Nova Technologies Limited"),
                             numberer.number("ORG_NAME", "NOVA"))
            self.assertNotEqual(numberer.number("ORG_NAME", "NOVA"),
                                numberer.number("ORG_NAME", "Zenith Labs Limited"))

            out = os.path.join(tmp, "sample_redacted.docx")
            R.redact_file(src, types=["ORG_NAME"], out_path=out, restore_mode=True)
            redacted = "\n".join(p.text for p in Document(out).paragraphs)
            placeholder = numberer.placeholder("ORG_NAME", "NOVA", orgs["NOVA"]["label"])
            self.assertEqual(redacted.count(placeholder), 4)
            self.assertNotIn("Nova Technologies Limited", redacted)
            self.assertNotIn("NOVA", redacted)
            self.assertTrue(os.path.isfile(out + ".mapping.json"))
            with ZipFile(out) as archive:
                self.assertNotIn("customXml/mapping.json", archive.namelist())

            restored = os.path.join(tmp, "sample_restored.docx")
            R.restore_file(out, out_path=restored)
            self.assertEqual("\n".join(p.text for p in Document(restored).paragraphs), SAMPLE)


if __name__ == "__main__":
    unittest.main()

"""Focused external-review workflow and OOXML safety checks on synthetic DOCX files."""
import io
import json
import os
import sys
import tempfile
import threading
import urllib.error
import urllib.request
import uuid
import zipfile

ROOT = os.path.dirname(os.path.dirname(__file__))
sys.path.insert(0, os.path.join(ROOT, "scripts") if os.path.isdir(os.path.join(ROOT, "scripts")) else ROOT)

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from lxml import etree
import contract_app_server as app
import contract_review_ext as review_ext


def request(base, path, data=None, ctype=None):
    headers = {"Content-Type": ctype} if ctype else {}
    req = urllib.request.Request(base + path, data=data, headers=headers,
                                 method="POST" if data is not None else "GET")
    try:
        with urllib.request.urlopen(req, timeout=30) as response:
            return json.loads(response.read()), response.status
    except urllib.error.HTTPError as exc:
        return json.loads(exc.read()), exc.code


def upload(base, path, case_name="演示案件"):
    boundary = "test" + uuid.uuid4().hex
    with open(path, "rb") as source:
        data = source.read()
    chunks = []
    for name, value in (("mode", "redact"), ("profile", "litigation"), ("case_name", case_name)):
        chunks.append((f"--{boundary}\r\nContent-Disposition: form-data; name=\"{name}\"\r\n\r\n{value}\r\n").encode())
    chunks.append((f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"{os.path.basename(path)}\"\r\n"
                   "Content-Type: application/vnd.openxmlformats-officedocument.wordprocessingml.document\r\n\r\n").encode())
    chunks.append(data)
    chunks.append((f"\r\n--{boundary}--\r\n").encode())
    return request(base, "/api/upload", b"".join(chunks), f"multipart/form-data; boundary={boundary}")


def make_doc(path, text, hidden=False):
    doc = Document()
    para = doc.add_paragraph()
    run = para.add_run(text)
    if hidden:
        run.font.hidden = True
    doc.save(path)


def run():
    with tempfile.TemporaryDirectory(prefix="external_review_") as tmp:
        server = app.start_app_server("127.0.0.1", 0, os.path.join(tmp, "sessions"))
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        base = f"http://127.0.0.1:{server.server_port}"
        try:
            first = os.path.join(tmp, "first.docx")
            make_doc(first, "案号：（2026）京0105民初123号 北京市朝阳区人民法院 证人张三 提到海星计划。")
            up, code = upload(base, first)
            assert code == 200 and up["ok"], up
            sid = up["sid"]
            initial = server.RequestHandlerClass.state.get(sid)["report"]
            types = {it["type"] for it in initial["items"]}
            assert {"CASE_NUMBER", "COURT_NAME"}.issubset(types), types

            added, code = request(base, f"/api/term/{sid}", json.dumps({"action": "add", "term": {
                "label": "项目代号", "values": ["海星计划", "海星"]}}).encode(), "application/json")
            assert code == 200 and added["ok"], added
            info = server.RequestHandlerClass.state.get(sid)
            review = info["review_payload"]
            assert any(m["label"] == "项目代号" for b in review["blocks"] for m in b["marks"])
            applied, code = request(base, f"/api/apply/{sid}", json.dumps({"enabled_types": [
                t["id"] for t in review["types"]]}).encode(), "application/json")
            assert code == 200 and applied["ok"], applied
            output = os.path.join(tmp, "sessions", sid, "out", applied["output_name"])
            with zipfile.ZipFile(output) as archive:
                names = archive.namelist()
                assert not any(n.startswith("customXml/") or n.startswith("docProps/") or "comments" in n for n in names), names
                assert all("海星" not in archive.read(n).decode("utf-8", "ignore") for n in names)
                root = etree.fromstring(archive.read("word/document.xml"))
                ignorable = root.get("{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable", "")
                assert all(prefix in root.nsmap for prefix in ignorable.split()), root.nsmap
            first_text = "".join(p.text for p in Document(output).paragraphs)
            assert "海星" not in first_text and "[项目代号#1]" in first_text, first_text
            assert os.path.exists(output + ".mapping.json")

            second = os.path.join(tmp, "second.docx")
            make_doc(second, "证人张三在海星会议上发言。")
            up2, code = upload(base, second)
            assert code == 200 and up2["ok"], up2
            info2 = server.RequestHandlerClass.state.get(up2["sid"])
            marks = [m for b in info2["review_payload"]["blocks"] for m in b["marks"]]
            assert any(m["placeholder"] == "[项目代号#1]" for m in marks), marks
            local, code = request(base, f"/api/term/{up2['sid']}", json.dumps({"action": "add", "scope": "document",
                "term": {"label": "内部词", "values": ["海星会议"]}}).encode(), "application/json")
            assert code == 200 and local["ok"], local
            up2b, code = upload(base, second)
            assert code == 200 and up2b["ok"], up2b
            info2b = server.RequestHandlerClass.state.get(up2b["sid"])
            assert not any(t["label"] == "内部词" for t in info2b["review_payload"]["types"])

            structured = os.path.join(tmp, "structured.docx")
            structured_doc = Document()
            structured_doc.sections[0].header.paragraphs[0].text = "秘密客户"
            structured_doc.add_table(rows=1, cols=1).cell(0, 0).text = "秘密客户"
            structured_doc.save(structured)
            up_struct, code = upload(base, structured, "结构案")
            assert code == 200 and up_struct["ok"], up_struct
            term_struct, code = request(base, f"/api/term/{up_struct['sid']}", json.dumps({"action": "add",
                "term": {"label": "客户", "values": ["秘密客户"]}}).encode(), "application/json")
            assert code == 200 and term_struct["ok"], term_struct
            info_struct = server.RequestHandlerClass.state.get(up_struct["sid"])
            exported, code = request(base, f"/api/apply/{up_struct['sid']}", json.dumps({"enabled_types": [
                t["id"] for t in info_struct["review_payload"]["types"]]}).encode(), "application/json")
            assert code == 200 and exported["ok"], exported
            structured_out = Document(os.path.join(tmp, "sessions", up_struct["sid"], "out", exported["output_name"]))
            assert "[客户#1]" in structured_out.sections[0].header.paragraphs[0].text
            assert "[客户#1]" in structured_out.tables[0].cell(0, 0).text

            hidden = os.path.join(tmp, "hidden.docx")
            make_doc(hidden, "证人张三，隐藏文字。", hidden=True)
            up3, code = upload(base, hidden, "另一案")
            assert code == 200 and up3["ok"], up3
            term3, code = request(base, f"/api/term/{up3['sid']}",
                                  json.dumps({"action": "add", "term": {"label": "秘密", "values": ["隐藏文字"]}}).encode(),
                                  "application/json")
            assert code == 200 and term3["ok"], term3
            info3 = server.RequestHandlerClass.state.get(up3["sid"])
            fail, status = request(base, f"/api/apply/{up3['sid']}",
                                   json.dumps({"enabled_types": [t["id"] for t in info3["review_payload"]["types"]]}).encode(),
                                   "application/json")
            assert status == 500 and not fail["ok"] and "隐藏" in fail["error"], fail

            pdf = os.path.join(tmp, "fake.pdf")
            with open(pdf, "wb") as f:
                f.write(b"%PDF-1.4")
            rejected, status = upload(base, pdf, "演示案件")
            assert status == 400 and "仅支持" in rejected["error"], rejected

            commented = os.path.join(tmp, "commented.docx")
            make_doc(commented, "已脱敏正文")
            with zipfile.ZipFile(commented, "a") as archive:
                archive.writestr("word/comments.xml", "<comments>批注里有原文秘密</comments>")
            with zipfile.ZipFile(commented) as archive:
                parts = {n: archive.read(n) for n in archive.namelist()}
            parts["word/document.xml"] = parts["word/document.xml"].replace(
                b"</w:p>", b'<w:commentRangeStart w:id="0"/></w:p>', 1)
            with zipfile.ZipFile(commented, "w") as archive:
                for name, data in parts.items():
                    archive.writestr(name, data)
            review_ext.clean_docx(commented, [])
            with zipfile.ZipFile(commented) as archive:
                assert not any("comments" in n or n.startswith("customXml/") for n in archive.namelist())
                assert b"commentRangeStart" not in archive.read("word/document.xml")

            revised = os.path.join(tmp, "revised.docx")
            make_doc(revised, "已脱敏正文")
            with zipfile.ZipFile(revised, "r") as archive:
                parts = {n: archive.read(n) for n in archive.namelist()}
            parts["word/document.xml"] = parts["word/document.xml"].replace(
                b"</w:body>", "<w:ins><w:r><w:t>修订残留</w:t></w:r></w:ins></w:body>".encode())
            with zipfile.ZipFile(revised, "w") as archive:
                for name, data in parts.items():
                    archive.writestr(name, data)
            try:
                review_ext.clean_docx(revised, [])
                raise AssertionError("修订内容应阻断外发")
            except ValueError as exc:
                assert "修订" in str(exc), exc

            hidden_style = os.path.join(tmp, "hidden-style.docx")
            styled = Document()
            secret_style = styled.styles.add_style("SecretStyle", WD_STYLE_TYPE.CHARACTER)
            secret_style.font.hidden = True
            styled.add_paragraph().add_run("样式隐藏的文字", style=secret_style)
            styled.save(hidden_style)
            try:
                review_ext.clean_docx(hidden_style, [])
                raise AssertionError("隐藏样式应阻断外发")
            except ValueError as exc:
                assert "隐藏" in str(exc), exc
            print("PASS: 自选字段、别名、同案代号、诉讼预设、干净 DOCX、隐藏/修订阻断、批注剥离、PDF 阻断")
        finally:
            server.shutdown()
            server.server_close()
            thread.join(timeout=2)


if __name__ == "__main__":
    run()
